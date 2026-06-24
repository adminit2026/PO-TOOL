"""
Excel file generation services for PO Review application
Contains helper functions to reduce code duplication in Excel generation
"""
import logging
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment


# Excel styling constants
HEADER_FILL = PatternFill(start_color="0A5F9C", end_color="0A5F9C", fill_type="solid")
HEADER_FONT = Font(bold=True, size=11, color="FFFFFF")
GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
RED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
YELLOW_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")


def auto_adjust_column_width(ws, max_width=50):
    """
    Auto-adjust column widths based on content
    
    Args:
        ws: Worksheet object
        max_width: Maximum column width
    """
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except (TypeError, AttributeError):
                pass
        adjusted_width = min(max_length + 2, max_width)
        ws.column_dimensions[column_letter].width = adjusted_width


def write_headers(ws, headers, row=1):
    """
    Write styled headers to worksheet
    
    Args:
        ws: Worksheet object
        headers: List of header strings
        row: Row number to write headers (default: 1)
    """
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.fill = HEADER_FILL


def is_field_edited(header, edited_fields):
    """
    Check if a specific field was edited
    
    Args:
        header: Column header name
        edited_fields: Dictionary of edited fields
        
    Returns:
        Boolean indicating if field was edited
    """
    field_map = {
        'Quantity': 'quantity',
        'Unit Cost': 'unitCost',
        'Production Cost': 'productionCost',
        'UPS Cost': 'upsCost',
        'Box Number': 'boxNumber'
    }
    field_key = field_map.get(header)
    return field_key and edited_fields.get(field_key, False)


def write_sheet_data(ws, headers, data, color_by_status=True):
    """
    Write data rows with conditional formatting based on approval status
    
    Args:
        ws: Worksheet object
        headers: List of column headers
        data: List of data dictionaries
        color_by_status: Apply color coding based on approval status
    """
    for row_idx, item in enumerate(data, 2):
        is_approved = item.get('Approval Status') == 'approved'
        is_rejected = item.get('Approval Status') == 'rejected'
        edited_fields = item.get('edited', {})
        
        for col_idx, header in enumerate(headers, 1):
            value = item.get(header, '')
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            
            if color_by_status:
                # Yellow for edited cells, Green for approved, Red for rejected
                if is_field_edited(header, edited_fields):
                    cell.fill = YELLOW_FILL
                elif is_approved:
                    cell.fill = GREEN_FILL
                elif is_rejected:
                    cell.fill = RED_FILL


def create_excel_with_approval(results_list: list, output_path: str):
    """
    Create Excel file with separate sheets for Approved, Rejected, and Pending orders
    with conditional color formatting
    
    Args:
        results_list: List of result dictionaries
        output_path: Path to save the Excel file
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    
    headers = [
        'PO', 'Vendor', 'Ship to Location', 'ASIN', 'External ID', 'Model Number', 'Title',
        'Quantity', 'Unit Cost', 'Production Cost', 'UPS Cost', 'Operational Cost', 
        'Commission', 'Total Cost/Unit', 'Margin/Unit', 'Margin %', 
        'Total Cost', 'Total Margin', 'Stock Quantity', 'Sales Units (30d)', 'Status'
    ]
    
    # Separate data by approval status
    approved_items = [item for item in results_list if item.get('Approval Status') == 'approved']
    rejected_items = [item for item in results_list if item.get('Approval Status') == 'rejected']
    pending_items = [item for item in results_list if item.get('Approval Status', 'pending') == 'pending']
    
    def create_sheet(name, data, color_status):
        """Helper to create a single sheet"""
        ws = wb.create_sheet(name)
        write_headers(ws, headers)
        write_sheet_data(ws, headers, data, color_by_status=color_status)
        auto_adjust_column_width(ws)
        return ws
    
    # Create All Orders sheet first
    ws_all = create_sheet("All Orders", results_list, True)
    wb.move_sheet(ws_all, offset=-len(wb.sheetnames) + 1)  # Move to beginning
    
    # Create other sheets
    if approved_items:
        create_sheet("Approved Orders", approved_items, True)
    if rejected_items:
        create_sheet("Rejected Orders", rejected_items, True)
    if pending_items:
        create_sheet("Pending Orders", pending_items, False)
    
    wb.save(output_path)


def create_export_file(approved_items: list, output_path: str):
    """
    Create EXPORT file with approved items
    Columns: PO, Vendor, Ship to location, Model Number, ASIN, External ID, Title, Availability
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "POout"
    
    headers = ['PO', 'Vendor', 'Ship to location', 'Model Number', 'ASIN', 'External ID', 'Title', 'Availability']
    write_headers(ws, headers)
    
    for row_idx, item in enumerate(approved_items, 2):
        ws.cell(row=row_idx, column=1, value=item.get('PO', ''))
        ws.cell(row=row_idx, column=2, value=item.get('Vendor', ''))
        ws.cell(row=row_idx, column=3, value=item.get('Ship to Location', ''))
        ws.cell(row=row_idx, column=4, value=item.get('Model Number', ''))
        ws.cell(row=row_idx, column=5, value=item.get('ASIN', ''))
        ws.cell(row=row_idx, column=6, value=item.get('External ID', ''))
        ws.cell(row=row_idx, column=7, value=item.get('Title', ''))
        ws.cell(row=row_idx, column=8, value='Accepted: In stock')
    
    auto_adjust_column_width(ws)
    wb.save(output_path)


def create_box_file(approved_items: list, output_path: str):
    """
    Create BOX file with box numbers
    Columns: PO, Vendor, Ship to location, Model Number, Expected Quantity, BOX
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "POout"
    
    headers = ['PO', 'Vendor', 'Ship to location', 'Model Number', 'Expected Quantity', 'BOX']
    write_headers(ws, headers)
    
    for row_idx, item in enumerate(approved_items, 2):
        ws.cell(row=row_idx, column=1, value=item.get('PO', ''))
        ws.cell(row=row_idx, column=2, value=item.get('Vendor', ''))
        ws.cell(row=row_idx, column=3, value=item.get('Ship to Location', ''))
        ws.cell(row=row_idx, column=4, value=item.get('Model Number', ''))
        ws.cell(row=row_idx, column=5, value=item.get('Quantity', 0))
        ws.cell(row=row_idx, column=6, value=item.get('Box Number', ''))
    
    auto_adjust_column_width(ws, max_width=30)
    wb.save(output_path)


def load_reference_data():
    """Load reference data for Color and Taille matching"""
    try:
        ref_file = '/app/backend/reference_data.xlsx'
        df = pd.read_excel(ref_file, engine='openpyxl')
        reference_dict = {}
        for _, row in df.iterrows():
            sku = row.get('SKU', '')
            if pd.notna(sku):
                reference_dict[str(sku).strip()] = {
                    'color': row.get('Color', ''),
                    'taille': row.get('Taille', ''),
                    'prod_or_stock': row.get('Prod or Stock', '')
                }
        return reference_dict
    except Exception as e:
        logging.error(f"Error loading reference data: {e}")
        return {}


def create_production_sheets(approved_items: list, output_path: str):
    """
    Create Production Sheets with two tabs:
    - 'To Print Stock' for SKUs starting with 'J'
    - 'To Print Prod' for all other SKUs
    """
    wb = openpyxl.Workbook()
    reference_dict = load_reference_data()
    
    # Add reference data to items
    items_with_ref = []
    for item in approved_items:
        model = item.get('Model Number', '')
        ref_data = reference_dict.get(str(model).strip(), {})
        item_with_ref = {
            **item,
            'Color': ref_data.get('color', '') if pd.notna(ref_data.get('color')) else '',
            'Taille': ref_data.get('taille', '') if pd.notna(ref_data.get('taille')) else ''
        }
        items_with_ref.append(item_with_ref)
    
    # Separate by SKU prefix
    stock_items = [item for item in items_with_ref if str(item.get('Model Number', '')).strip().upper().startswith('J')]
    prod_items = [item for item in items_with_ref if not str(item.get('Model Number', '')).strip().upper().startswith('J')]
    
    headers = ['REF', 'COLOR', 'TAILLE', 'QTY']
    
    def write_production_sheet(ws, items):
        """Helper to write production sheet data"""
        write_headers(ws, headers)
        for row_idx, item in enumerate(items, 2):
            ws.cell(row=row_idx, column=1, value=item.get('Model Number', ''))
            ws.cell(row=row_idx, column=2, value=item.get('Color', ''))
            ws.cell(row=row_idx, column=3, value=item.get('Taille', ''))
            ws.cell(row=row_idx, column=4, value=item.get('Quantity', 0))
        auto_adjust_column_width(ws)
    
    # Create sheets
    if prod_items:
        ws_prod = wb.active
        ws_prod.title = 'To Print Prod'
        write_production_sheet(ws_prod, prod_items)
    
    if stock_items:
        ws_stock = wb.create_sheet('To Print Stock', 0)
        write_production_sheet(ws_stock, stock_items)
    
    # Remove default sheet if no prod items
    if not prod_items and 'Sheet' in wb.sheetnames:
        wb.remove(wb['Sheet'])
    
    wb.save(output_path)


def create_ean_list_csv_by_location(approved_items: list, output_dir: str):
    """
    Create separate EAN List CSV files for each location and return ZIP file
    
    Format: semicolon-separated (EAN;ASIN;SKU;QTY) for ALL locations
    
    Args:
        approved_items: List of approved item dictionaries
        output_dir: Directory to save the ZIP file
        
    Returns:
        Path to ZIP file containing all CSVs
    """
    import zipfile
    from pathlib import Path
    
    # Group items by location
    location_groups = {}
    for item in approved_items:
        location = item.get('Ship to Location', 'Unknown')
        location_code = location.split('-')[0].strip() if '-' in location else location.strip()
        
        if location_code not in location_groups:
            location_groups[location_code] = []
        location_groups[location_code].append(item)
    
    # Create directory for CSV files
    csv_dir = Path(output_dir) / 'ean_csvs'
    csv_dir.mkdir(parents=True, exist_ok=True)
    
    csv_files = []
    
    # Create CSV for each location - ALL use semicolon format
    for location_code, items in location_groups.items():
        filename = f"{location_code} EAN.csv"
        filepath = csv_dir / filename
        
        with open(filepath, 'w', newline='', encoding='utf-8') as csvfile:
            # Format: EAN;ASIN;SKU;QTY
            for item in items:
                line = f"{item.get('External ID', '')};{item.get('ASIN', '')};{item.get('Model Number', '')};{item.get('Quantity', 0)}\n"
                csvfile.write(line)
        
        csv_files.append(filepath)
    
    # Create ZIP file
    zip_path = Path(output_dir) / 'EAN_Lists.zip'
    with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as zipf:
        for csv_file in csv_files:
            zipf.write(csv_file, csv_file.name)
    
    return str(zip_path)


def create_packing_list(approved_items: list, output_path: str):
    """
    Create Packing List Excel
    Columns: Destination (Location), Reference (SKU), QTY, Carton AMZNCC/SSCC
    
    Args:
        approved_items: List of approved item dictionaries
        output_path: Path to save the Excel file
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Packing List"
    
    headers = ['Destination', 'Reference', 'QTY', 'Carton AMZNCC/SSCC']
    write_headers(ws, headers)
    
    # Write data
    for row_idx, item in enumerate(approved_items, 2):
        ws.cell(row=row_idx, column=1, value=item.get('Ship to Location', ''))
        ws.cell(row=row_idx, column=2, value=item.get('Model Number', ''))
        ws.cell(row=row_idx, column=3, value=item.get('Quantity', 0))
        ws.cell(row=row_idx, column=4, value=item.get('Box Number', ''))
    
    auto_adjust_column_width(ws, max_width=40)
    wb.save(output_path)
